# Copyright (c) Kulix  
# See the file 'LICENSE' for your permissions while using this content 
# ----------------------------------------------------------------------------------------------------------------------------------------------------------|   
#     - Do not touch or modify the code below. If there is an error, please contact the owner, but under no circumstances should you touch the code  
#     - In the case you wish to analyze this code under any circumstance; note that touching this code should not be altered under any circumstance  
#            - If this code is altered, the owner holds no obligation for damages or errors caused by the altered code  
#     - Do not resell this tool, do not credit it as your own

import sys
import os

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))

from Core.Config.Info import *
from Core.Config.Utilities import *

try:
    from PIL import Image
    from PIL.ExifTags import TAGS
    import PyPDF2
    import mutagen
    import time
    import openpyxl
    from pptx import Presentation
    from tkinter import filedialog
    import tkinter as tk
    from docx import Document
except Exception as e:
    ErrorModule(e)

Title("Metadata Extractor")

def ExtractImageMetadata(file_path):
    try:
        image = Image.open(file_path)
        exif_data = image._getexif()
        metadata = {}
        if exif_data:
            for tag, value in exif_data.items():
                tag_name = TAGS.get(tag, tag)
                metadata[tag_name] = value
        return metadata
    except Exception as e:
        ErrorModule(e)
        return None

def ExtractPDFMetadata(file_path):
    try:
        with open(file_path, "rb") as file:
            reader = PyPDF2.PdfReader(file)
            metadata = reader.metadata
            return metadata
    except Exception as e:
        ErrorModule(e)
        return None

def ExtractAudioMetadata(file_path):
    try:
        audio = mutagen.File(file_path)
        return audio.tags if audio else None
    except Exception as e:
        ErrorModule(e)
        return None

def ExtractWordMetadata(file_path):
    try:
        doc = Document(file_path)
        metadata = {}
        core_props = doc.core_properties
        metadata["Title"] = core_props.title
        metadata["Author"] = core_props.author
        metadata["Subject"] = core_props.subject
        metadata["Keywords"] = core_props.keywords
        metadata["Last Modified By"] = core_props.last_modified_by
        metadata["Created"] = core_props.created
        metadata["Modified"] = core_props.modified
        return metadata
    except Exception as e:
        ErrorModule(e)
        return None

def ExtractExcelMetadata(file_path):
    try:
        wb = openpyxl.load_workbook(file_path)
        metadata = {}
        properties = wb.properties
        metadata["Title"] = properties.title
        metadata["Creator"] = properties.creator
        metadata["Last Modified By"] = properties.last_modified_by
        metadata["Created"] = properties.created
        metadata["Modified"] = properties.modified
        return metadata
    except Exception as e:
        ErrorModule(e)
        return None

def ExtractPowerPointMetadata(file_path):
    try:
        ppt = Presentation(file_path)
        metadata = {}
        properties = ppt.core_properties
        metadata["Title"] = properties.title
        metadata["Author"] = properties.author
        metadata["Subject"] = properties.subject
        metadata["Keywords"] = properties.keywords
        metadata["Last Modified By"] = properties.last_modified_by
        metadata["Created"] = properties.created
        metadata["Modified"] = properties.modified
        return metadata
    except Exception as e:
        ErrorModule(e)
        return None

def ExtractTextMetadata(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            metadata = {"File Size": os.path.getsize(file_path)}
        return metadata
    except Exception as e:
        ErrorModule(e)
        return None

def ExtractMetadata(file_path):
    print(f"\n{BEFORE_CYAN + current_time_hour() + AFTER_CYAN} {GEN_VALID} Extracting metadata from: {green}{file_path}")
    
    ext = os.path.splitext(file_path)[1].lower()
    metadata = None
    
    if ext in [".jpg", ".jpeg", ".png", ".tiff"]:
        metadata = ExtractImageMetadata(file_path)
    elif ext in [".pdf"]:
        metadata = ExtractPDFMetadata(file_path)
    elif ext in [".mp3", ".flac", ".wav", ".ogg"]:
        metadata = ExtractAudioMetadata(file_path)
    elif ext in [".docx", ".doc"]:
        metadata = ExtractWordMetadata(file_path)
    elif ext in [".xlsx", ".xls"]:
        metadata = ExtractExcelMetadata(file_path)
    elif ext in [".pptx", ".ppt"]:
        metadata = ExtractPowerPointMetadata(file_path)
    elif ext in [".txt"]:
        metadata = ExtractTextMetadata(file_path)
    else:
        print(f"{BEFORE_CYAN + current_time_hour() + AFTER_CYAN} {ERROR} Unsupported file type: {white}{ext}{cyan}")
        return
    
    if metadata:
        if not any(metadata.values()):
            print(f"{BEFORE_CYAN + current_time_hour() + AFTER_CYAN} {ERROR} Metadata on the file is empty or has been wiped previously{cyan}")
        else:
            for key, value in metadata.items():
                print(f"{BEFORE_CYAN + current_time_hour() + AFTER_CYAN} {ADD} {key}: {white}{value}{cyan}")
    else:
        print(f"{BEFORE_CYAN + current_time_hour() + AFTER_CYAN} {ERROR} Metadata on the file is empty or has been wiped previously{cyan}")

def SelectFile():
    root = tk.Tk()
    root.withdraw()
    file_path = filedialog.askopenfilename(title="Select a file", filetypes=(("All Files", "*.*"),))
    return file_path

try:
    file_path = SelectFile()
    if file_path:
        ExtractMetadata(file_path)
    else:
        print(f"{BEFORE_CYAN + current_time_hour() + AFTER_CYAN} {ERROR} No file selected.{cyan}")
    Continue()
    Reset()

except Exception as e:
    ErrorModule(e)
